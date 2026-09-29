import re
import struct
import zlib

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches

from core.language_tool_client import Match


class FakeLanguageTool:
    """Offline stand-in for LanguageToolClient driven by regex rules."""

    DEFAULT_RULES = [
        (r"\bresults shows\b", "results show", "GRAMMAR", "grammar", "AGREEMENT"),
        (r"\bemissions was\b", "emissions were", "GRAMMAR", "grammar", "AGREEMENT"),
        (r"\brecieved\b", "received", "TYPOS", "misspelling", "SPELL"),
        (r"\bteh\b", "the", "TYPOS", "misspelling", "SPELL"),
        (r"\bTeh\b", "The", "TYPOS", "misspelling", "SPELL"),
        (r" \.", ".", "PUNCTUATION", "typographical", "SPACE_BEFORE_DOT"),
        (r"\bin order to\b", "to", "REDUNDANCY", "style", "IN_ORDER_TO"),
        (r" ,", ",", "PUNCTUATION", "typographical", "COMMA_SPACE"),
        (r"\bGhawanmehh\b", "Ghawanmeh", "TYPOS", "misspelling", "SPELL"),
        (r"\bSimaPro\b", "Simper", "TYPOS", "misspelling", "SPELL"),
    ]

    def __init__(self, rules=None, fail=False):
        self.rules = rules or self.DEFAULT_RULES
        self.fail = fail
        self.request_count = 0
        self.texts = []

    def check_matches(self, text):
        from core.language_tool_client import LanguageToolError
        self.request_count += 1
        self.texts.append(text)
        if self.fail:
            raise LanguageToolError("service unavailable")
        matches = []
        for pattern, replacement, category, issue, rule in self.rules:
            for m in re.finditer(pattern, text):
                matches.append(Match(m.start(), m.end(), [replacement], rule, category, issue, f"{rule} message"))
        return matches


@pytest.fixture
def fake_lt():
    return FakeLanguageTool()


def tiny_png():
    raw = b"\x00\xff\x00\x00"
    def chunk(kind, data):
        return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data))
    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0))
            + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b""))


def add_hyperlink(paragraph, url, text):
    part = paragraph.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                          is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    text_el = OxmlElement("w:t")
    text_el.text = text
    text_el.set(qn("xml:space"), "preserve")
    run.append(text_el)
    link.append(run)
    paragraph._p.append(link)
    return link


@pytest.fixture
def sample_docx(tmp_path):
    doc = Document()
    doc.add_paragraph("Emission Study", style="Title")
    doc.add_heading("Introductoin", level=1)  # heading typo must survive by default

    p = doc.add_paragraph("The ")
    p.add_run("results shows").bold = True
    p.add_run(" that emissions was increased significantly in 2023 (Smith et al., 2024).")

    p = doc.add_paragraph("Data was recieved from ")
    add_hyperlink(p, "https://www.who.int/", "teh WHO portal")
    p.add_run(" in order to model CO")
    sub = p.add_run("2")
    sub.font.subscript = True
    p.add_run(" with SimaPro .")

    doc.add_paragraph("The author Ghawanmehh wrote this sentence.")

    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Teh fuel recieved"
    table.cell(0, 1).text = "1,200"
    table.cell(1, 0).text = "Diesel"
    table.cell(1, 1).text = "2,500 t"

    doc.add_paragraph().add_run().add_picture(_png_file(tmp_path), width=Inches(0.5))

    doc.sections[0].header.paragraphs[0].text = "Teh University header"

    doc.add_heading("References", level=1)
    doc.add_paragraph("Smith, J. (2024). Teh results shows emissions. Journal, 1(2), 3-4.")
    doc.add_heading("Appendix", level=1)
    doc.add_paragraph("Teh appendix text.")

    path = tmp_path / "sample.docx"
    doc.save(path)
    return path


def _png_file(tmp_path):
    path = tmp_path / "pixel.png"
    path.write_bytes(tiny_png())
    return str(path)
