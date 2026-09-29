from docx import Document

from core.document_reader import paragraph_segments, paragraph_text
from core.document_writer import Edit, apply_edits, apply_edits_to_paragraph, diff_edits, write_segments
from core.options import ProcessingOptions
from core.paragraph_processor import ParagraphProcessor
from core.pipeline import DocumentProcessor
from core.structure_manager import Kind, analyze_document
from tests.conftest import FakeLanguageTool, add_hyperlink


def make_paragraph(*runs):
    doc = Document()
    p = doc.add_paragraph()
    for text, bold in runs:
        p.add_run(text).bold = bold
    return doc, p


def runs(p):
    return [(r.text, r.bold) for r in p.runs]


def test_edit_inside_bold_run_keeps_bold():
    _, p = make_paragraph(("The ", None), ("results shows", True), (" that.", None))
    apply_edits_to_paragraph(p, [Edit(4, 17, "results show")])
    assert runs(p) == [("The ", None), ("results show", True), (" that.", None)]


def test_edit_spanning_runs_uses_first_run_format():
    _, p = make_paragraph(("Data wa", None), ("s recieved", True), (" today.", None))
    apply_edits_to_paragraph(p, [Edit(5, 17, "were received")])
    assert p.text == "Data were received today."
    assert runs(p)[0] == ("Data were received", None)
    assert runs(p)[1] == ("", True)


def test_insertion_and_multiple_edits():
    _, p = make_paragraph(("A quick fox", None), (" jumps", True))
    apply_edits_to_paragraph(p, [Edit(2, 2, "very "), Edit(11, 17, " leaps")])
    assert p.text == "A very quick fox leaps"


def test_edits_never_touch_hyperlinks():
    doc = Document()
    p = doc.add_paragraph("Visit ")
    add_hyperlink(p, "https://example.com", "teh site")
    p.add_run(" now.")
    segments = paragraph_segments(p)
    assert paragraph_text(segments) == "Visit teh site now."
    applied, skipped = apply_edits(segments, [Edit(6, 9, "the"), Edit(0, 5, "See")])
    write_segments(segments)
    assert [e.replacement for e in skipped] == ["the"]
    assert paragraph_text(paragraph_segments(p)) == "See teh site now."
    assert len(p._p.xpath(".//w:hyperlink")) == 1


def test_atomic_mode_is_all_or_nothing():
    doc = Document()
    p = doc.add_paragraph("Visit ")
    add_hyperlink(p, "https://example.com", "site")
    segments = paragraph_segments(p)
    applied, skipped = apply_edits(segments, [Edit(0, 5, "See"), Edit(6, 10, "page")], atomic=True)
    assert applied == [] and len(skipped) == 2


def test_diff_edits_reconstructs_text():
    old = "The results shows that emissions was increased (Smith, 2024)."
    new = "Results show that emissions increased (Smith, 2024)."
    doc, p = make_paragraph((old[:12], None), (old[12:30], True), (old[30:], None))
    apply_edits_to_paragraph(p, diff_edits(old, new), atomic=True)
    assert p.text == new


def test_paragraph_processor_simple_api():
    _, p = make_paragraph(("Teh results shows ", None), ("clearly", True), (".", None))
    result = ParagraphProcessor(FakeLanguageTool()).process(p)
    assert result == "The results show clearly."
    assert runs(p)[1] == ("clearly", True)


def test_structure_classification(sample_docx):
    units = analyze_document(Document(sample_docx))
    by_text = {u.paragraph.text: u for u in units}
    assert by_text["Emission Study"].kind == Kind.TITLE
    assert by_text["Introductoin"].kind == Kind.HEADING
    assert by_text["Smith, J. (2024). Teh results shows emissions. Journal, 1(2), 3-4."].kind == Kind.REFERENCE
    assert by_text["Teh appendix text."].kind == Kind.NORMAL
    assert by_text["1,200"].numeric_cell
    assert any(u.location.value == "header" for u in units)


def test_full_pipeline(sample_docx, tmp_path):
    out = tmp_path / "out.docx"
    report = DocumentProcessor(FakeLanguageTool(), options=ProcessingOptions()).process(sample_docx, out)
    doc = Document(out)
    texts = [p.text for p in doc.paragraphs]

    assert "The results show that emissions were increased significantly in 2023 (Smith et al., 2024)." in texts
    bold = [r for r in doc.paragraphs[2].runs if r.bold]
    assert [r.text for r in bold] == ["results show"]
    # hyperlink text, subscript and protected term untouched; other fixes applied
    assert "Data was received from teh WHO portal to model CO2 with SimaPro." in texts
    assert "The author Ghawanmehh wrote this sentence." in texts            # proper noun kept
    assert "Introductoin" in texts                                         # headings skipped
    assert "Smith, J. (2024). Teh results shows emissions. Journal, 1(2), 3-4." in texts  # references
    assert "The appendix text." in texts
    assert doc.tables[0].cell(0, 0).text == "The fuel received"
    assert doc.sections[0].header.paragraphs[0].text == "Teh University header"
    assert report.ok, report.to_text()
    assert report.corrections == 8
    assert report.protected["citation"] == 1
    assert report.document["images"] == 1


def test_headers_and_headings_when_enabled(sample_docx, tmp_path):
    out = tmp_path / "out.docx"
    opts = ProcessingOptions(process_headers=True, process_headings=True)
    DocumentProcessor(FakeLanguageTool(), options=opts).process(sample_docx, out)
    assert Document(out).sections[0].header.paragraphs[0].text == "The University header"


def test_category_toggles(sample_docx, tmp_path):
    out = tmp_path / "out.docx"
    opts = ProcessingOptions(spelling=False, style=False)
    report = DocumentProcessor(FakeLanguageTool(), options=opts).process(sample_docx, out)
    assert set(report.corrections_by_group) <= {"grammar", "punctuation"}
    assert "Teh appendix text." in [p.text for p in Document(out).paragraphs]


def test_languagetool_failure_keeps_original(sample_docx, tmp_path):
    out = tmp_path / "out.docx"
    report = DocumentProcessor(FakeLanguageTool(fail=True), options=ProcessingOptions()).process(sample_docx, out)
    assert report.errors and report.corrections == 0
    assert "Teh appendix text." in [p.text for p in Document(out).paragraphs]


def test_small_batches_give_same_result(sample_docx, tmp_path):
    a, b = tmp_path / "a.docx", tmp_path / "b.docx"
    lt = FakeLanguageTool()
    DocumentProcessor(FakeLanguageTool(), options=ProcessingOptions(), max_chars=1000).process(sample_docx, a)
    DocumentProcessor(lt, options=ProcessingOptions(), max_chars=40).process(sample_docx, b)
    assert lt.request_count > 1
    assert [p.text for p in Document(a).paragraphs] == [p.text for p in Document(b).paragraphs]


def test_case_only_spelling_suggestions_are_rejected():
    from core.language_tool_client import Match
    from core.paragraph_processor import MatchFilter
    from core.protection import ProtectionEngine
    options = ProcessingOptions()
    f = MatchFilter(options, ProtectionEngine(options))
    text = "Consumption (kt) of fuel"
    match = Match(13, 15, ["KT"], "MORFOLOGIK_RULE_EN_US", "TYPOS", "misspelling", "")
    assert f.accept(match, text, []) == (False, "case-only suggestion")
