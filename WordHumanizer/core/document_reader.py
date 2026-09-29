"""Reading Word documents at run level.

A paragraph is turned into a list of *segments*: one per run, in document
order, each knowing its text, its position inside the paragraph's text and
whether it may be edited. Runs are not editable when changing them could
damage the document: hyperlinks, field results (cross-references, Zotero /
Mendeley / EndNote citations), tracked changes, content controls, runs that
hold images or other objects, hidden text, and superscript/subscript runs
(footnote-style citations, chemical formulas such as CO₂).
"""

from dataclasses import dataclass

from docx import Document
from docx.oxml.ns import qn
from docx.text.run import Run

MATH_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"

_EDITABLE_RUN_CHILDREN = {qn("w:rPr"), qn("w:t"), qn("w:tab"), qn("w:br"), qn("w:cr"),
                          qn("w:lastRenderedPageBreak")}
_LOCKED_CONTAINERS = {
    qn("w:hyperlink"): "hyperlink",
    qn("w:ins"): "tracked-change",
    qn("w:moveTo"): "tracked-change",
    qn("w:fldSimple"): "field",
    qn("w:sdt"): "content-control",
    qn("w:smartTag"): "smart-tag",
    qn("w:customXml"): "custom-xml",
}
_SKIPPED = {qn("w:pPr"), qn("w:del"), qn("w:moveFrom")}


@dataclass
class Segment:
    run: object          # docx.text.run.Run or None for a barrier
    text: str
    start: int
    editable: bool
    reason: str = ""     # why the segment is locked
    modified: bool = False

    @property
    def end(self):
        return self.start + len(self.text)


class WordReader:

    def __init__(self, filename):
        self.filename = filename
        self.document = Document(filename)

    def paragraphs(self):
        return self.document.paragraphs

    def tables(self):
        return self.document.tables

    def sections(self):
        return self.document.sections


def paragraph_segments(paragraph):
    """Return the list of Segments of a python-docx Paragraph."""
    segments = []
    state = {"field_depth": 0, "pos": 0}
    _collect(paragraph._p, paragraph, segments, state, locked_reason="")
    return segments


def paragraph_text(segments):
    return "".join(s.text for s in segments)


def has_math(paragraph):
    return bool(paragraph._p.xpath(".//m:oMath | .//m:oMathPara"))


def _collect(element, paragraph, segments, state, locked_reason):
    for child in element.iterchildren():
        tag = child.tag
        if tag in _SKIPPED or not isinstance(tag, str):
            continue
        if tag == qn("w:r"):
            _add_run(child, paragraph, segments, state, locked_reason)
        elif tag in _LOCKED_CONTAINERS:
            _collect(_content_of(child), paragraph, segments, state,
                     locked_reason or _LOCKED_CONTAINERS[tag])
        elif tag.startswith("{%s}" % MATH_NS):
            # Equations: zero-width barrier so no edit can span across them.
            segments.append(Segment(None, "", state["pos"], False, "equation"))
        # bookmarks, proofErr, comment ranges, etc. carry no text


def _content_of(element):
    if element.tag == qn("w:sdt"):
        content = element.find(qn("w:sdtContent"))
        return content if content is not None else element
    return element


def _add_run(r, paragraph, segments, state, locked_reason):
    run = Run(r, paragraph)
    # Complex fields: runs between fldChar begin/end are field codes/results.
    field_chars = r.findall(qn("w:fldChar"))
    in_field_before = state["field_depth"] > 0
    for fc in field_chars:
        kind = fc.get(qn("w:fldCharType"))
        if kind == "begin":
            state["field_depth"] += 1
        elif kind == "end":
            state["field_depth"] = max(0, state["field_depth"] - 1)

    text = run.text
    reason = locked_reason or _run_lock_reason(r)
    if not reason and (in_field_before or field_chars or state["field_depth"] > 0):
        reason = "field"
    segments.append(Segment(run, text, state["pos"], editable=not reason, reason=reason))
    state["pos"] += len(text)


def _run_lock_reason(r):
    for child in r.iterchildren():
        if child.tag not in _EDITABLE_RUN_CHILDREN:
            return "object"
        if child.tag == qn("w:br") and child.get(qn("w:type")) not in (None, "textWrapping"):
            return "page-break"
    rpr = r.find(qn("w:rPr"))
    if rpr is not None:
        valign = rpr.find(qn("w:vertAlign"))
        if valign is not None and valign.get(qn("w:val")) in ("superscript", "subscript"):
            return "super/subscript"
        if rpr.find(qn("w:vanish")) is not None:
            return "hidden"
        style = rpr.find(qn("w:rStyle"))
        if style is not None and "hyperlink" in (style.get(qn("w:val")) or "").lower():
            return "hyperlink"
    return ""
