from types import SimpleNamespace

import pytest
from docx import Document

from core.options import ProcessingOptions
from core.pipeline import DocumentProcessor
from core.rewriting_engine import (ClaudeRewritingEngine, NoOpRewritingEngine, RewriteError,
                                   RewritingEngine, create_rewriting_engine)
from tests.conftest import FakeLanguageTool, add_hyperlink


class ScriptedRewriter(RewritingEngine):
    available = True

    def __init__(self, func):
        self.func = func
        self.seen = []

    def rewrite(self, text, style="natural", language="en-US"):
        self.seen.append(text)
        return self.func(text)


def build(tmp_path, *paragraphs):
    doc = Document()
    for runs in paragraphs:
        p = doc.add_paragraph()
        for text, bold in runs:
            p.add_run(text).bold = bold
    path = tmp_path / "in.docx"
    doc.save(path)
    return path


SENTENCE = "It was found by the researchers that emissions increased by 12% in 2023 (Smith et al., 2024)."


def run(tmp_path, rewriter, path=None, **opts):
    path = path or build(tmp_path, [(SENTENCE[:30], None), (SENTENCE[30:], True)])
    out = tmp_path / "out.docx"
    options = ProcessingOptions(rewrite=True, **opts)
    report = DocumentProcessor(FakeLanguageTool(rules=[(r"^$", "", "GRAMMAR", "grammar", "X")]),
                               rewriter, options).process(path, out)
    return report, Document(out)


def test_rewrite_applied_and_protected_content_masked(tmp_path):
    def rewrite(masked):
        assert "12%" not in masked and "Smith" not in masked and "2023" not in masked
        return masked.replace("It was found by the researchers that emissions", "The researchers found that emissions")
    rewriter = ScriptedRewriter(rewrite)
    report, doc = run(tmp_path, rewriter)
    assert doc.paragraphs[0].text == ("The researchers found that emissions increased by 12% in 2023 "
                                      "(Smith et al., 2024).")
    assert report.rewritten == 1 and report.rewrites_rejected == 0
    assert any(r.bold for r in doc.paragraphs[0].runs)


def test_rewrite_that_drops_placeholder_is_rejected(tmp_path):
    report, doc = run(tmp_path, ScriptedRewriter(lambda m: m.split("⟦")[0] + "."))
    assert doc.paragraphs[0].text == SENTENCE
    assert report.rewrites_rejected == 1 and report.rewritten == 0


def test_rewrite_that_changes_numbers_is_rejected(tmp_path):
    # model writes a literal number instead of keeping the placeholder
    def rewrite(masked):
        return masked.replace("⟦P0⟧", "15%", 1)
    report, doc = run(tmp_path, ScriptedRewriter(rewrite))
    assert doc.paragraphs[0].text == SENTENCE
    assert report.rewrites_rejected == 1


def test_rewrite_touching_hyperlink_is_rejected(tmp_path):
    doc = Document()
    p = doc.add_paragraph("According to the official portal of the organisation located at ")
    add_hyperlink(p, "https://who.int", "the WHO website")
    p.add_run(" the data were published last year for everyone.")
    path = tmp_path / "link.docx"
    doc.save(path)
    # moves the placeholder for the locked hyperlink run to the end
    def rewrite(masked):
        import re
        placeholders = re.findall(r"⟦P\d+⟧", masked)
        body = re.sub(r"\s*⟦P\d+⟧\s*", " ", masked).strip()
        return body + " " + " ".join(placeholders)
    report, out = run(tmp_path, ScriptedRewriter(rewrite), path=path)
    assert out.paragraphs[0].text == ("According to the official portal of the organisation located at "
                                      "the WHO website the data were published last year for everyone.")
    assert report.rewrites_rejected == 1


def test_rewriter_errors_are_isolated(tmp_path):
    def boom(masked):
        raise RewriteError("declined")
    report, doc = run(tmp_path, ScriptedRewriter(boom))
    assert doc.paragraphs[0].text == SENTENCE and report.rewrites_rejected == 1


def test_short_paragraphs_and_headings_not_rewritten(tmp_path):
    rewriter = ScriptedRewriter(lambda m: m.upper())
    path = build(tmp_path, [("Too short to rewrite.", None)])
    report, _ = run(tmp_path, rewriter, path=path)
    assert rewriter.seen == []


def test_rewrite_without_provider_warns(tmp_path):
    report, _ = run(tmp_path, NoOpRewritingEngine())
    assert report.rewritten == 0 and any("REWRITE_PROVIDER" in w for w in report.warnings)


def test_claude_engine_request_and_refusal():
    calls = []

    def create(**kwargs):
        calls.append(kwargs)
        text = SimpleNamespace(type="text", text="  Better text.  ")
        return SimpleNamespace(stop_reason=calls and kwargs.get("_stop", "end_turn"), content=[text])

    client = SimpleNamespace(beta=SimpleNamespace(messages=SimpleNamespace(create=create)))
    engine = ClaudeRewritingEngine(model="claude-opus-5-5", client=client)
    assert engine.rewrite("Some ⟦P0⟧ text.", style="academic") == "Better text."
    kwargs = calls[0]
    assert kwargs["model"] == "claude-opus-5-5"
    assert kwargs["thinking"] == {"type": "adaptive"}
    assert "⟦P0⟧" in kwargs["system"] and "academic" in kwargs["messages"][0]["content"]

    refusal = SimpleNamespace(beta=SimpleNamespace(messages=SimpleNamespace(
        create=lambda **k: SimpleNamespace(stop_reason="refusal", content=[]))))
    with pytest.raises(RewriteError):
        ClaudeRewritingEngine(client=refusal).rewrite("x")


def test_factory_defaults_to_noop():
    settings = SimpleNamespace(rewrite_provider="none", rewrite_model="m")
    assert isinstance(create_rewriting_engine(settings), NoOpRewritingEngine)
